from pptx import Presentation
from pptx.util import Inches
from pptx.oxml import parse_xml

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Transition XML
xml_str = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:push dir="r"/></p:transition>'
trans_elem = parse_xml(xml_str)
slide._element.append(trans_elem)

prs.save("test_trans.pptx")
print("Transition test successful!")
